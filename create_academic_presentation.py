"""
AgriYield - Full Academic Review Presentation Generator (8 Slides)
Covers all requirements:
1. Cover Slide (S. SANDHIYA - ROLL NO : 24VNM034)
2. Synopsis of Project
3. Module Description
4. DFD (Data Flow Diagram)
5. ERD (Entity Relationship Diagram)
6. Table Design (Database Schemas)
7. Form Design (Input Screen UI)
8. Output Screen (Results UI)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = os.path.join(os.path.dirname(__file__), "AgriYield_Presentation.pptx")

# Light Theme Palette
BG_COLOR = RGBColor(248, 250, 252)       # Soft off-white #f8fafc
WHITE = RGBColor(255, 255, 255)          # Card background #ffffff
CARD_BORDER = RGBColor(226, 232, 240)    # Soft gray border #e2e8f0
CARD_BG_GREEN = RGBColor(240, 253, 244)  # Light mint tint #f0fdf4
CARD_BG_BLUE = RGBColor(240, 249, 255)   # Light blue tint #f0f9ff

PRIMARY_DARK = RGBColor(15, 23, 42)      # Deep slate for main text #0f172a
BODY_MUTED = RGBColor(51, 65, 85)        # Readable dark slate #334155
EMERALD = RGBColor(5, 150, 105)          # Forest emerald #059669
EMERALD_DARK = RGBColor(6, 95, 70)       # Deep forest #065f46
ACCENT_BLUE = RGBColor(2, 132, 199)      # Sky blue #0284c7
ACCENT_PURPLE = RGBColor(124, 58, 237)   # Purple #7c3aed


def set_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_slide_header(slide, slide_num, title, total_slides=8):
    # Top Tag
    tag_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.35), Inches(11.5), Inches(0.3))
    p_tag = tag_box.text_frame.paragraphs[0]
    p_tag.text = f"AGRIYIELD  •  SLIDE {slide_num} OF {total_slides}"
    p_tag.font.size = Pt(10.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = EMERALD

    # Slide Title
    t_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.65), Inches(11.5), Inches(0.75))
    p_title = t_box.text_frame.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = PRIMARY_DARK


def build_academic_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ========================================================
    # SLIDE 1: COVER SLIDE
    # ========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(0.8), Inches(11.533), Inches(0.12))
    banner.fill.solid()
    banner.fill.fore_color.rgb = EMERALD
    banner.line.color.rgb = EMERALD

    t_box = s1.shapes.add_textbox(Inches(0.9), Inches(1.15), Inches(11.533), Inches(1.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "AgriYield: Data Science-Based Crop Yield\nAnalysis and Prediction"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD_DARK

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Academic Project Review  •  Machine Learning & Web Dashboard Interface"
    p1_sub.font.size = Pt(16)
    p1_sub.font.color.rgb = BODY_MUTED
    p1_sub.space_before = Pt(6)

    # Presenter Card
    auth_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.3), Inches(6.0), Inches(3.6))
    auth_card.fill.solid()
    auth_card.fill.fore_color.rgb = CARD_BG_GREEN
    auth_card.line.color.rgb = EMERALD
    auth_card.line.width = Pt(2.5)

    at_box = s1.shapes.add_textbox(Inches(1.2), Inches(3.55), Inches(5.4), Inches(3.1))
    at_tf = at_box.text_frame
    at_tf.word_wrap = True
    
    p_a1 = at_tf.paragraphs[0]
    p_a1.text = "PROJECT PRESENTED BY"
    p_a1.font.size = Pt(13)
    p_a1.font.bold = True
    p_a1.font.color.rgb = EMERALD

    p_a2 = at_tf.add_paragraph()
    p_a2.text = "S. SANDHIYA"
    p_a2.font.size = Pt(34)
    p_a2.font.bold = True
    p_a2.font.color.rgb = PRIMARY_DARK
    p_a2.space_before = Pt(8)

    p_a3 = at_tf.add_paragraph()
    p_a3.text = "ROLL NO : 24VNM034"
    p_a3.font.size = Pt(20)
    p_a3.font.bold = True
    p_a3.font.color.rgb = ACCENT_BLUE
    p_a3.space_before = Pt(8)

    p_a4 = at_tf.add_paragraph()
    p_a4.text = "Department of Computer Science & Engineering"
    p_a4.font.size = Pt(14)
    p_a4.font.color.rgb = BODY_MUTED
    p_a4.space_before = Pt(8)

    # Highlights Card
    info_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(3.3), Inches(5.2), Inches(3.6))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = WHITE
    info_card.line.color.rgb = CARD_BORDER
    info_card.line.width = Pt(1.5)

    it_box = s1.shapes.add_textbox(Inches(7.5), Inches(3.55), Inches(4.6), Inches(3.1))
    it_tf = it_box.text_frame
    it_tf.word_wrap = True
    
    p_i1 = it_tf.paragraphs[0]
    p_i1.text = "PROJECT HIGHLIGHTS"
    p_i1.font.size = Pt(13)
    p_i1.font.bold = True
    p_i1.font.color.rgb = EMERALD

    facts = [
        "Dataset Scope: 19,689 Records (1997–2020)",
        "Coverage: 55 Crops across 30 Indian States",
        "Champion Model: Random Forest Regressor",
        "Holdout Accuracy: 96.8% (R² Score)",
        "Zero Data Leakage: Genuine Pre-Harvest ML"
    ]
    for f in facts:
        p = it_tf.add_paragraph()
        p.text = f"•  {f}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK
        p.space_before = Pt(8)

    # ========================================================
    # SLIDE 2: SYNOPSIS OF PROJECT
    # ========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_slide_header(s2, 2, "Synopsis of Project")

    syn_cards = [
        ("PROJECT ABSTRACT",
         "AgriYield is a machine learning data science system designed to accurately forecast crop yield prior to harvest. "
         "Using 24 years of longitudinal agricultural census records across India, the system captures non-linear relationships "
         "between crop species, soil-climate factors, and agricultural inputs to produce reliable harvest estimates.",
         EMERALD),
        ("PROBLEM STATEMENT",
         "Traditional agricultural yield estimation relies on manual crop-cutting experiments conducted after harvest. "
         "This results in delayed reporting, preventing timely food security interventions, logistics coordination, and market price stabilization.",
         ACCENT_BLUE),
        ("OBJECTIVES & SCOPE",
         "• Predict crop yield (Tonnes/Ha) months before harvest using pre-sowing factors.\n"
         "• Prevent target data leakage by strictly excluding post-harvest production.\n"
         "• Deploy a farmer-friendly web dashboard providing acre yields and truckload counts.",
         ACCENT_PURPLE)
    ]

    for idx, (title, text, color) in enumerate(syn_cards):
        x = Inches(0.9 + idx * 3.9)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(5.4))
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE
        c.line.color.rgb = color
        c.line.width = Pt(2)

        tb = s2.shapes.add_textbox(x + Inches(0.25), Inches(1.75), Inches(3.2), Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_b = tf.add_paragraph()
        p_b.text = text
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = PRIMARY_DARK
        p_b.space_before = Pt(14)

    # ========================================================
    # SLIDE 3: MODULE DESCRIPTION
    # ========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_slide_header(s3, 3, "Module Description")

    modules = [
        ("Module 1: Data Ingestion & Audit",
         "Loads 19,689 census records from Ministry of Agriculture & IMD. Validates schema, strips categorical whitespaces, handles zero yields, and isolates Coconut unit scales.",
         EMERALD),
        ("Module 2: Feature Engineering & Leakage Defense",
         "Calculates Fertilizer/Area and Pesticide/Area rates. Strictly eliminates 'Production' to prevent mathematical leakage. Applies robust log-transforms.",
         ACCENT_BLUE),
        ("Module 3: Model Training & Evaluation",
         "Executes temporal train/test split (1997–2015 train vs 2016–2020 test). Benchmarks Ridge, Decision Tree, Gradient Boosting, and Random Forest.",
         ACCENT_PURPLE),
        ("Module 4: Prediction & Agronomic Inference",
         "Accepts user inputs, passes through the serialized pipeline, and outputs predicted yield, 90% confidence weather intervals, and regional benchmarks.",
         RGBColor(217, 119, 6)),
        ("Module 5: Executive Web Dashboard",
         "Streamlit presentation layer with quick scenarios, dynamic metric cards, Plotly comparison charts, and plain-language harvest summaries.",
         RGBColor(13, 148, 136))
    ]

    for idx, (m_title, m_desc, m_color) in enumerate(modules):
        y = Inches(1.5 + idx * 1.08)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), y, Inches(11.533), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = m_color
        card.line.width = Pt(1.5)

        tb = s3.shapes.add_textbox(Inches(1.1), y + Inches(0.1), Inches(11.1), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = m_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = m_color

        p_d = tf.add_paragraph()
        p_d.text = m_desc
        p_d.font.size = Pt(12.5)
        p_d.font.color.rgb = PRIMARY_DARK
        p_d.space_before = Pt(3)

    # ========================================================
    # SLIDE 4: DATA FLOW DIAGRAM (DFD)
    # ========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_slide_header(s4, 4, "Data Flow Diagram (DFD) — Level 1 Architecture")

    # Add diagram image
    if os.path.exists("dfd_diagram.png"):
        s4.shapes.add_picture("dfd_diagram.png", Inches(0.9), Inches(1.5), width=Inches(7.8))

    # Right side explanation card
    dfd_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(1.5), Inches(3.433), Inches(5.4))
    dfd_card.fill.solid()
    dfd_card.fill.fore_color.rgb = WHITE
    dfd_card.line.color.rgb = EMERALD
    dfd_card.line.width = Pt(2)

    tb_dfd = s4.shapes.add_textbox(Inches(9.2), Inches(1.75), Inches(3.0), Inches(4.8))
    tf_dfd = tb_dfd.text_frame
    tf_dfd.word_wrap = True

    p = tf_dfd.paragraphs[0]
    p.text = "DFD FLOW SUMMARY"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    dfd_pts = [
        ("User / Officer: ", "Enters crop, state, season, area, and inputs."),
        ("Process 1.0: ", "Validates domain bounds and ensures non-negativity."),
        ("Process 2.0: ", "Computes application rates and scales features."),
        ("Process 3.0: ", "Random Forest queries model weights (D2) and predicts."),
        ("Process 4.0: ", "Translates raw yield into acre and truckload units."),
        ("Stores D1 & D2: ", "Census database and pre-trained model artifacts.")
    ]
    for t, d in dfd_pts:
        p = tf_dfd.add_paragraph()
        p.text = f"• {t}{d}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = PRIMARY_DARK
        p.space_before = Pt(10)

    # ========================================================
    # SLIDE 5: ENTITY RELATIONSHIP DIAGRAM (ERD)
    # ========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_slide_header(s5, 5, "Entity-Relationship Diagram (ERD)")

    if os.path.exists("erd_diagram.png"):
        s5.shapes.add_picture("erd_diagram.png", Inches(0.9), Inches(1.5), width=Inches(7.8))

    erd_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(1.5), Inches(3.433), Inches(5.4))
    erd_card.fill.solid()
    erd_card.fill.fore_color.rgb = WHITE
    erd_card.line.color.rgb = ACCENT_BLUE
    erd_card.line.width = Pt(2)

    tb_erd = s5.shapes.add_textbox(Inches(9.2), Inches(1.75), Inches(3.0), Inches(4.8))
    tf_erd = tb_erd.text_frame
    tf_erd.word_wrap = True

    p = tf_erd.paragraphs[0]
    p.text = "ERD ENTITIES & SCHEMA"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    erd_pts = [
        ("STATE Entity: ", "Contains State_ID (PK), Name, and historical average yield."),
        ("CROP Entity: ", "Contains Crop_ID (PK), species name, category, and scale."),
        ("SEASON Entity: ", "Defines cropping calendar (Kharif, Rabi, Whole Year)."),
        ("CULTIVATION_RECORD: ", "Central entity storing area, rainfall, fertilizer, and pesticide."),
        ("YIELD_PREDICTION: ", "Output entity storing point predictions, acre rates, and truckloads."),
        ("Cardinality: ", "1:M relations connect master entities to cultivation events.")
    ]
    for t, d in erd_pts:
        p = tf_erd.add_paragraph()
        p.text = f"• {t}{d}"
        p.font.size = Pt(12.5)
        p.font.color.rgb = PRIMARY_DARK
        p.space_before = Pt(10)

    # ========================================================
    # SLIDE 6: TABLE DESIGN (DATABASE / SCHEMA)
    # ========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_slide_header(s6, 6, "Table Design & Relational Schema")

    # Table 1: CROP_CENSUS_RECORDS
    t1_box = s6.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.5), Inches(0.4))
    p1 = t1_box.text_frame.paragraphs[0]
    p1.text = "TABLE 1: CROP_CENSUS_RECORDS (Core Census Series — 19,689 Rows)"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = EMERALD_DARK

    t1_shape = s6.shapes.add_table(6, 4, Inches(0.9), Inches(1.8), Inches(11.533), Inches(1.9))
    tbl1 = t1_shape.table
    tbl1.columns[0].width = Inches(2.8)
    tbl1.columns[1].width = Inches(2.2)
    tbl1.columns[2].width = Inches(1.8)
    tbl1.columns[3].width = Inches(4.733)

    t1_headers = ["Field Name", "Data Type", "Constraints", "Description"]
    for i, h in enumerate(t1_headers):
        cell = tbl1.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = EMERALD_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE

    t1_rows = [
        ["Record_ID", "INTEGER", "PRIMARY KEY", "Unique auto-increment identifier"],
        ["Crop", "VARCHAR(50)", "NOT NULL", "Standardized crop species (e.g. Wheat, Rice, Sugarcane)"],
        ["State", "VARCHAR(50)", "NOT NULL", "Indian state or union territory administrative region"],
        ["Season", "VARCHAR(30)", "NOT NULL", "Cropping window (Kharif, Rabi, Summer, Whole Year)"],
        ["Area / Rainfall", "DOUBLE PRECISION", "CHECK (>= 0)", "Cultivated hectares (ha) and annual precipitation (mm)"]
    ]
    for r_idx, row in enumerate(t1_rows):
        for c_idx, val in enumerate(row):
            cell = tbl1.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11.5)
            p.font.color.rgb = PRIMARY_DARK

    # Table 2: YIELD_PREDICTIONS
    t2_box = s6.shapes.add_textbox(Inches(0.9), Inches(3.9), Inches(11.5), Inches(0.4))
    p2 = t2_box.text_frame.paragraphs[0]
    p2.text = "TABLE 2: YIELD_PREDICTIONS_LOG (Inference Output Schema)"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_BLUE

    t2_shape = s6.shapes.add_table(5, 4, Inches(0.9), Inches(4.3), Inches(11.533), Inches(1.6))
    tbl2 = t2_shape.table
    tbl2.columns[0].width = Inches(2.8)
    tbl2.columns[1].width = Inches(2.2)
    tbl2.columns[2].width = Inches(1.8)
    tbl2.columns[3].width = Inches(4.733)

    for i, h in enumerate(t1_headers):
        cell = tbl2.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE

    t2_rows = [
        ["Prediction_ID", "INTEGER", "PRIMARY KEY", "Unique prediction transaction identifier"],
        ["Predicted_Yield", "DOUBLE PRECISION", "NOT NULL", "Predicted agricultural rate in Tonnes per Hectare (t/ha)"],
        ["Yield_Per_Acre", "DOUBLE PRECISION", "NOT NULL", "Acre equivalent conversion (Predicted_Yield / 2.471)"],
        ["Total_Commercial_Tonnes", "DOUBLE PRECISION", "NOT NULL", "Gross harvest weight for entire farm (Predicted_Yield * Area)"]
    ]
    for r_idx, row in enumerate(t2_rows):
        for c_idx, val in enumerate(row):
            cell = tbl2.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11.5)
            p.font.color.rgb = PRIMARY_DARK

    take_box = s6.shapes.add_textbox(Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.9))
    tf_tk = take_box.text_frame
    p_tk = tf_tk.paragraphs[0]
    p_tk.text = "🛡️ Schema Integrity Note: 'Production' is explicitly omitted from training schemas to enforce zero data leakage."
    p_tk.font.size = Pt(12)
    p_tk.font.bold = True
    p_tk.font.color.rgb = EMERALD_DARK

    # ========================================================
    # SLIDE 7: FORM DESIGN (INPUT INTERFACE)
    # ========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_slide_header(s7, 7, "Form Design (User Input Interface)")

    # Real Form Screenshot
    if os.path.exists("form_design_screenshot.png"):
        s7.shapes.add_picture("form_design_screenshot.png", Inches(0.9), Inches(1.5), width=Inches(7.2))

    # Right Explanation Card
    form_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.5), Inches(4.133), Inches(5.4))
    form_card.fill.solid()
    form_card.fill.fore_color.rgb = WHITE
    form_card.line.color.rgb = EMERALD
    form_card.line.width = Pt(2)

    tb_fm = s7.shapes.add_textbox(Inches(8.55), Inches(1.75), Inches(3.6), Inches(4.8))
    tf_fm = tb_fm.text_frame
    tf_fm.word_wrap = True

    p = tf_fm.paragraphs[0]
    p.text = "INPUT FORM SPECIFICATIONS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    fm_pts = [
        ("Quick Demonstration Presets: ", "1-Click buttons for Punjab Wheat, West Bengal Rice, and Maharashtra Sugarcane."),
        ("Crop & State Dropdowns: ", "Select from 55 authentic crops and 30 Indian states."),
        ("Cropping Season: ", "Select from Kharif, Rabi, Summer, Autumn, Winter, or Whole Year."),
        ("Cultivated Area: ", "Numeric input in Hectares (district or farm acreage)."),
        ("Rainfall & Inputs: ", "Annual precipitation (mm), total fertilizer (kg), and pesticide (kg)."),
        ("Action Trigger: ", "Prominent green 'Calculate Predicted Yield' execution button.")
    ]
    for t, d in fm_pts:
        p = tf_fm.add_paragraph()
        p.text = f"• {t}{d}"
        p.font.size = Pt(12)
        p.font.color.rgb = PRIMARY_DARK
        p.space_before = Pt(8)

    # ========================================================
    # SLIDE 8: OUTPUT SCREEN (RESULTS INTERFACE)
    # ========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_slide_header(s8, 8, "Output Screen (Prediction Results)")

    # Real Output Screenshot
    if os.path.exists("output_screen_screenshot.png"):
        s8.shapes.add_picture("output_screen_screenshot.png", Inches(0.9), Inches(1.5), width=Inches(7.2))

    # Right Explanation Card
    out_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.5), Inches(4.133), Inches(5.4))
    out_card.fill.solid()
    out_card.fill.fore_color.rgb = WHITE
    out_card.line.color.rgb = ACCENT_BLUE
    out_card.line.width = Pt(2)

    tb_out = s8.shapes.add_textbox(Inches(8.55), Inches(1.75), Inches(3.6), Inches(4.8))
    tf_out = tb_out.text_frame
    tf_out.word_wrap = True

    p = tf_out.paragraphs[0]
    p.text = "OUTPUT SCREEN COMPONENTS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    out_pts = [
        ("Simple Summary Box: ", "Provides a plain-language harvest summary stating total tonnes and acre equivalents."),
        ("Card 1 — Harvest Rate: ", "Displays yield rate (e.g. 3.00 t/ha and 1.21 tonnes/acre) with state productivity tier badge."),
        ("Card 2 — Total Production: ", "Displays gross farm tonnage (30,000 tonnes) and commercial 10-tonne truckload counts (3,000 trucks)."),
        ("Card 3 — Weather Sensitivity: ", "90% confidence range (2.5 – 3.5 t/ha) accounting for dry vs. favorable monsoon seasons."),
        ("Benchmark Comparison: ", "Interactive bar chart comparing farm prediction with state average and national average.")
    ]
    for t, d in out_pts:
        p = tf_out.add_paragraph()
        p.text = f"• {t}{d}"
        p.font.size = Pt(12)
        p.font.color.rgb = PRIMARY_DARK
        p.space_before = Pt(8)

    prs.save(PPTX_PATH)
    print("Academic 8-slide presentation generated successfully at:", PPTX_PATH)


if __name__ == "__main__":
    build_academic_presentation()
